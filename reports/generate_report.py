from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

doc = Document()

# Base font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

# Title
title = doc.add_heading('Rockmould Limited', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = ACCENT

subtitle = doc.add_paragraph('Website & Platform Update Report')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(16)
subtitle.runs[0].font.bold = True

date_p = doc.add_paragraph(f"Date: {date.today().strftime('%B %d, %Y')}")
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.runs[0].font.size = Pt(11)
date_p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()

# Intro
intro = doc.add_paragraph(
    "This report summarizes the improvements and fixes completed today across the "
    "Rockmould website, admin/CMS dashboard, and client portal. The work covered "
    "visual design, usability, responsiveness across devices, and behind-the-scenes "
    "code cleanup and quality checks."
)

# Section helper
def add_item(heading_text, problem_text, fix_text):
    h = doc.add_heading(heading_text, level=2)
    for run in h.runs:
        run.font.color.rgb = ACCENT

    p1 = doc.add_paragraph()
    r1 = p1.add_run('Issue: ')
    r1.bold = True
    p1.add_run(problem_text)

    p2 = doc.add_paragraph()
    r2 = p2.add_run('Resolution: ')
    r2.bold = True
    p2.add_run(fix_text)

# ============ SECTION 1: EARLIER THIS EVENING ============
h1 = doc.add_heading('Part 1 — Site-Wide Responsiveness & Layout Improvements', level=1)
for run in h1.runs:
    run.font.color.rgb = ACCENT

add_item(
    "1. Responsive Cards, Calendar & Project Listings",
    "The luxury information cards, the consultation booking calendar, and the project "
    "listing pages were not adapting well between mobile and desktop screens. Some "
    "content was overflowing or not resizing as expected.",
    "Spacing, sizing, and layout rules were adjusted across approximately 18 admin and "
    "public-facing pages so these elements now resize and reflow correctly on phones, "
    "tablets, and desktop computers."
)

add_item(
    "2. Duplicate Header Controls on Desktop",
    "On desktop screens, certain navigation/header controls were appearing twice, "
    "creating a cluttered and unprofessional look.",
    "The duplication was removed so the header now displays a single, clean set of "
    "controls."
)

add_item(
    "3. Removal of Unused Default Page",
    "The platform still contained a leftover demo \"Welcome\" page from the underlying "
    "software framework, which is no longer used since the site has its own homepage.",
    "This unused file was removed to keep the codebase clean and reduce clutter."
)

add_item(
    "4. Hamburger Navigation Menu Restored on Desktop",
    "After the responsiveness improvements, the hamburger (☰) menu — which is the "
    "only way visitors can open the site navigation, as there is no separate desktop "
    "menu bar — had become hidden on desktop screens.",
    "The hamburger menu was restored and now stays visible on desktop alongside the "
    "\"Client Portal\" and \"Get in Touch\" buttons, ensuring visitors can always access "
    "the navigation."
)

add_item(
    "5. Better Use of Space on Large Monitors",
    "On large desktop displays, laptops with high resolutions, and ultrawide monitors, "
    "the website content was capped at a fixed width, leaving large empty margins on "
    "either side.",
    "Additional sizing rules were added so the content area expands proportionally on "
    "larger screens (1536px, 1920px, and 2560px and above), making better use of the "
    "available screen space on big monitors."
)

# ============ SECTION 2: THIS SESSION ============
h2 = doc.add_heading('Part 2 — Admin/CMS, Portal & Code Quality Improvements', level=1)
for run in h2.runs:
    run.font.color.rgb = ACCENT

add_item(
    "1. Oversized Headings in Admin & Client Portal",
    "Page titles in the admin dashboard and client portal (e.g. \"Consultations\", "
    "\"Projects\", \"CMS Homepage\") were displaying far too large — closer to a "
    "marketing banner than a normal page heading.",
    "The heading styles were corrected so titles now display at a clear, readable size "
    "across all admin and portal pages. This was checked and confirmed on the "
    "Consultations, Dashboard, Projects, KYC, CMS Homepage, CMS Services, Payments, "
    "Documents, Updates, and Media pages."
)

add_item(
    "2. Cluttered Consultation Booking Calendar",
    "On the public consultation booking page, each calendar day was showing both a "
    "date number and a repeated weekday label (e.g. \"15\" and \"Mon\" together), even "
    "though the column headers above already display the weekday names.",
    "The redundant weekday label was removed from each day cell. The calendar is now "
    "cleaner and easier to read, particularly on mobile devices."
)

add_item(
    "3. Outdated Code Cleanup",
    "The codebase contained four old layout template files left over from a previous "
    "restructuring, which were no longer used anywhere in the system. Unused duplicate "
    "files increase the risk of future changes being made to the wrong file.",
    "These four unused files were safely removed. The full automated test suite "
    "(54 tests) was run afterward to confirm that nothing was broken — all tests "
    "passed."
)

add_item(
    "4. Backend Error Check",
    "A review of the application's error logs was carried out to check for any active "
    "issues affecting the platform.",
    "No active errors were found. One historical error from a previously resolved bug "
    "was found in older logs, but it does not affect the current system and required "
    "no further action."
)

add_item(
    "5. Full Responsive Design Review (Mobile to Desktop)",
    "A review was carried out to confirm that the public website, admin dashboard, and "
    "client portal all adapt correctly across phone, tablet, and desktop screen sizes "
    "— including tables, grids, and navigation menus.",
    "The review confirmed that tables scroll correctly on small screens, page layouts "
    "stack appropriately on mobile, and both the public site and admin/portal areas "
    "have working mobile navigation menus. No further changes were required."
)

# ============ SUMMARY ============
h3 = doc.add_heading('Summary', level=1)
for run in h3.runs:
    run.font.color.rgb = ACCENT

doc.add_paragraph(
    "All planned improvement items for today have been completed, covering: light/dark "
    "theme consistency, typography sizing, visual polish, removal of outdated code, "
    "hero text sizing, project card alignment, the consultation booking flow, and "
    "responsiveness across all device sizes."
)

h4 = doc.add_heading('Outstanding Action for Client', level=2)
for run in h4.runs:
    run.font.color.rgb = ACCENT

doc.add_paragraph(
    "For security purposes, please change the admin panel login password from the "
    "current default to a new, private password at your earliest convenience."
)

doc.save(r"C:\xampp\htdocs\rockmould-legacy-main\reports\Rockmould_Update_Report.docx")
print("done")
