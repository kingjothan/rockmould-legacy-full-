from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x66, 0x66, 0x66)

title = doc.add_heading('Rockmould Limited', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = ACCENT

subtitle = doc.add_paragraph('Website & Platform — Full System Overview')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(16)
subtitle.runs[0].font.bold = True

date_p = doc.add_paragraph(f"Date: {date.today().strftime('%B %d, %Y')}")
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_p.runs[0].font.size = Pt(11)
date_p.runs[0].font.color.rgb = GREY

doc.add_paragraph()

intro = doc.add_paragraph(
    "This document explains, in plain language, everything the Rockmould Limited "
    "website and online platform currently does. It covers the public website that "
    "visitors see, the consultation booking system, the secure Client Portal used by "
    "investors/clients, and the Admin Dashboard used by Rockmould staff to manage "
    "everything behind the scenes. The goal is to give a clear picture of how the "
    "system works as a whole, the workflows that connect each part, and what each "
    "screen or feature is for."
)

# helper for level-1 heading
def h1(text):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = ACCENT
    return h

# helper for level-2 heading
def h2(text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = ACCENT
    return h

# helper for a bullet list
def bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

# helper for a labelled paragraph
def labelled(label, text):
    p = doc.add_paragraph()
    r = p.add_run(label + ': ')
    r.bold = True
    p.add_run(text)


# ============================================================
# SECTION 1: PUBLIC WEBSITE
# ============================================================
h1('1. The Public Website')

doc.add_paragraph(
    "This is the part of the site anyone on the internet can visit. It presents "
    "Rockmould Limited to the public, showcases projects and services, and lets "
    "visitors get in touch, apply for jobs, book consultations, and sign up for "
    "updates."
)

h2('1.1 Homepage')
bullets([
    "A hero banner with the company's main message and call-to-action buttons.",
    "A showcase of featured projects.",
    "An overview of the services Rockmould offers.",
    "A 'Why Choose Us' section listing key strengths and benefits.",
    "A step-by-step overview of how Rockmould works with clients (the process).",
    "Client testimonials.",
    "A preview of recent blog articles/insights.",
    "A call-to-action pointing visitors to the Client Portal.",
    "Company statistics (e.g. projects completed, years in business, clients served).",
])
doc.add_paragraph(
    "Every section of the homepage can be edited by Rockmould staff through the "
    "Admin CMS — no developer is needed to change the wording, images, or featured "
    "items."
)

h2('1.2 About Page')
bullets([
    "Company story and description.",
    "A list of team members, each with a photo and short bio.",
    "A company timeline showing key milestones in Rockmould's history.",
    "Key company metrics (projects completed, years active, clients, professionals).",
])

h2('1.3 Services Page')
bullets([
    "A list of all services Rockmould offers, each with a description, benefits, "
    "and the steps involved.",
    "Each service can be linked to related projects, so visitors can see real "
    "examples of that service in action.",
])

h2('1.4 Projects Directory & Project Pages')
bullets([
    "A searchable, filterable directory of all published projects (filter by "
    "category/type, status, or keyword search).",
    "Featured projects are highlighted at the top.",
    "Each individual project page shows: full description, key dates, milestones "
    "achieved, project updates, a media gallery (organised into sections such as "
    "exterior, interior, and general views), downloadable documents, related "
    "projects, and testimonials about that project.",
    "Visitors can submit an enquiry directly from a project page (Project Inquiry "
    "form), which is sent straight to the Rockmould team.",
])

h2('1.5 Portfolio Page')
bullets([
    "A curated showcase of Rockmould's best/portfolio-worthy projects, paired with "
    "testimonials and overall company metrics — used as a 'highlight reel' for new "
    "visitors.",
])

h2('1.6 Blog')
bullets([
    "A searchable, filterable list of articles/insights, organised by category.",
    "A featured article is shown prominently, with recent posts listed alongside.",
    "Each article has its own page with related articles suggested at the bottom.",
])

h2('1.7 FAQ Page')
bullets([
    "Frequently asked questions, organised into categories, with a search box so "
    "visitors can quickly find answers.",
])

h2('1.8 Careers Page')
bullets([
    "A list of currently open job positions.",
    "Each job has its own page with full details and an application form. "
    "Applications are emailed to the Rockmould team and stored for follow-up.",
])

h2('1.9 Contact, Partnerships & Testimonials Pages')
bullets([
    "Contact page: provides the contact form visitors use to message Rockmould "
    "directly (name, email, phone, department, message).",
    "Partnerships page: showcases featured projects and partnership opportunities, "
    "with its own enquiry form for prospective partners.",
    "Testimonials page: a full page dedicated to client reviews and feedback.",
])

h2('1.10 Gallery, Privacy, Terms & Utility Pages')
bullets([
    "Gallery: a media library of images and videos that can be filtered by type "
    "(image/video), project, or gallery section.",
    "Privacy Policy and Terms & Conditions: standard legal pages.",
    "Coming Soon and Maintenance pages: placeholder pages the site can switch to "
    "during planned downtime or before a new feature launches.",
])

h2('1.11 Forms Visitors Can Submit')
doc.add_paragraph(
    "All of the following are captured automatically, stored in the system, and "
    "trigger an email notification to the Rockmould team:"
)
bullets([
    "Contact Form — general enquiries.",
    "Newsletter Signup — adds the visitor's email to the mailing list.",
    "Partnership Enquiry — interest from prospective business partners.",
    "Project Enquiry — a question about a specific project, sent from that "
    "project's page.",
    "Career Application — a job application for a specific open position.",
])


# ============================================================
# SECTION 2: CONSULTATION BOOKING SYSTEM
# ============================================================
h1('2. Consultation Booking System')

doc.add_paragraph(
    "This feature lets website visitors book a one-on-one consultation with the "
    "Rockmould team directly from the website, without needing to call or email "
    "first."
)

h2('2.1 How a Visitor Books a Consultation')
bullets([
    "The visitor goes to the Consultation page and is shown an interactive "
    "calendar.",
    "Only dates and time slots that Rockmould has marked as available are "
    "selectable — fully booked days, blocked dates (e.g. holidays), and past "
    "dates are automatically greyed out.",
    "The visitor selects a consultation type, a date, a time slot, and enters a "
    "subject and their contact details.",
    "On submission, the system creates a booking with a unique reference number, "
    "shows a confirmation page, and emails a confirmation to both the visitor and "
    "the Rockmould team.",
])

h2('2.2 How Rockmould Staff Manage Consultations')
bullets([
    "Admin staff see a list of all bookings, with search and filters (by name, "
    "email, phone, type, subject, or status), plus quick stats: total, pending, "
    "confirmed, upcoming, and completed bookings.",
    "Staff can open a booking to assign a staff member, add a meeting link, and "
    "change its status (pending, confirmed, completed, cancelled, or no-show). The "
    "client is emailed automatically when the booking is updated.",
    "Staff configure the available consultation types, the timezone used, and "
    "which email address receives new-booking notifications.",
    "Staff set up a recurring weekly availability schedule (which days and times "
    "consultations can be booked) and can block off specific dates or date ranges "
    "(with a reason) when they are unavailable.",
])


# ============================================================
# SECTION 3: CLIENT PORTAL
# ============================================================
h1('3. Client Portal')

doc.add_paragraph(
    "The Client Portal is a private, secure area where Rockmould's clients "
    "(property investors/owners) can log in to track their projects, documents, "
    "media, payments, and communicate with Rockmould directly."
)

h2('3.1 Logging In (No Passwords — Email + One-Time Code)')
bullets([
    "Existing clients log in by entering their email address. The system checks "
    "that an active client account exists for that email, then sends a 6-digit "
    "one-time code (OTP) to that email address.",
    "The client enters the code to log in. The code expires after a short time and "
    "can only be tried a limited number of times (5 attempts) before it is locked, "
    "for security.",
    "If the code email doesn't arrive, the client can request a new one, but only "
    "once every 60 seconds — this prevents abuse.",
    "New clients (not yet registered) fill in a short verification form instead — "
    "see 'New Client Onboarding' below.",
])

h2('3.2 New Client Onboarding (KYC)')
bullets([
    "A prospective client who doesn't yet have portal access fills in a "
    "verification ('KYC' — Know Your Client) form: full name, phone number, email, "
    "residential address, the project they're interested in (or a custom answer if "
    "it's not listed), and optionally an identification document.",
    "This creates a pending request that Rockmould staff review and either approve "
    "or reject (see Admin section below). Once approved, the client can log in "
    "using the OTP process above.",
])

h2('3.3 Dashboard (Home Screen)')
bullets([
    "A summary of the client's investments at a glance: how many projects/units "
    "they own, how many assignments are currently active, how much they have paid "
    "in total, and how much balance is still outstanding.",
    "A quick view of the most recent documents, media, and project updates "
    "relevant to them.",
])

h2('3.4 My Projects')
bullets([
    "A list of every project/unit the client owns, with status.",
    "Opening a project shows: full project information, the specific unit "
    "assigned to them, all milestones reached so far, all progress updates marked "
    "visible to clients, all payments made for that unit, plus photos/videos and "
    "documents specific to that project.",
])

h2('3.5 Documents')
bullets([
    "A central library of all documents the client has been given access to "
    "(e.g. contracts, receipts, floor plans), searchable by title, category, or "
    "project.",
    "Clients can download any document they have access to.",
])

h2('3.6 Media Gallery')
bullets([
    "Photos and videos related to the client's own projects, filterable by type "
    "(image/video) or by project, with a preview/lightbox view.",
    "Clients can download media files.",
])

h2('3.7 Balance & Payments')
bullets([
    "A clear financial summary for each project/unit: the total contract amount, "
    "how much has been paid so far, and what remains outstanding.",
    "A full payment history per project, plus an overall total across all of the "
    "client's projects.",
])

h2('3.8 Project Updates')
bullets([
    "A timeline/feed of progress updates posted by Rockmould staff for the "
    "client's projects — each with a date, title, description, and a progress "
    "percentage — so clients can follow construction/delivery progress without "
    "needing to ask.",
])

h2('3.9 Profile')
bullets([
    "Clients can view their account details and update their email address, "
    "address, and a short bio.",
])

h2('3.10 Support')
bullets([
    "Clients can raise a support ticket — optionally linked to a specific "
    "project — with a subject and message, which goes directly to the Rockmould "
    "team.",
    "Clients can see the status of all tickets they have raised (open or closed).",
])


# ============================================================
# SECTION 4: ADMIN DASHBOARD
# ============================================================
h1('4. Admin Dashboard (Staff Side)')

doc.add_paragraph(
    "This is the control centre used by Rockmould staff to run the website, manage "
    "clients and projects, track finances, and respond to enquiries. Access is "
    "protected by a staff login (email + password), and different staff members can "
    "be given different levels of access depending on their role."
)

h2('4.1 Staff Roles & Permissions')
bullets([
    "Admin — full access to everything.",
    "Editor — manages website content (CMS).",
    "Project Manager — manages projects, client assignments, and consultations.",
    "Finance — manages payments and financial records.",
    "Content Manager — manages media, project updates, and website content.",
    "Only Admins can manage staff accounts/roles, KYC approvals, and system "
    "settings; other roles see only the sections relevant to their job.",
])

h2('4.2 Dashboard (Home Screen)')
bullets([
    "Key numbers at a glance: total active clients, total active project "
    "assignments, pending KYC requests, and outstanding balances owed by clients.",
    "An alerts panel flagging anything needing attention: pending consultations, "
    "pending KYC reviews, overdue payments, open support tickets, and new contact "
    "form submissions.",
    "A recent activity feed showing the latest actions taken by staff across the "
    "system.",
])

h2('4.3 Client Management')
bullets([
    "A searchable list of all clients (by name, email, or phone), filterable by "
    "status (active/inactive/pending).",
    "Each client's profile shows their personal details, all projects/units "
    "assigned to them, recent login (OTP) attempts for support purposes, their "
    "support tickets, and their documents.",
    "Staff can manually create a client account, edit a client's details, or "
    "delete a client (with a confirmation step). Deleting a client automatically "
    "frees up any units that were assigned to them.",
])

h2('4.4 KYC Review (New Client Approvals)')
bullets([
    "A list of all KYC submissions from the portal, filterable by status "
    "(pending/approved/rejected).",
    "Staff can open a request to view the applicant's details and identification "
    "document, then approve or reject it with an optional note.",
    "Approving a request automatically creates the client's portal account, ready "
    "for them to log in. Rejected applicants can be told why and can resubmit.",
    "Every decision is logged, including who reviewed it and when.",
])

h2('4.5 Consultation Management')
doc.add_paragraph(
    "Covered in Section 2.2 above — this is where staff manage bookings, "
    "availability, blocked dates, and consultation settings."
)

h2('4.6 Project Management & Client Assignments')
bullets([
    "A searchable list of all projects with the number of units and assignments "
    "for each.",
    "Each project page shows everything about that project: description, "
    "category, status, location, key dates, features/highlights, floor plans, "
    "documents, gallery, units, client assignments, milestones, and updates.",
    "Staff can create new projects, edit existing ones (including uploading "
    "images/media), and mark them as featured, portfolio, or publicly visible.",
    "Deleting a project requires typing the exact project title to confirm — "
    "preventing accidental deletion.",
])
h2('4.6.1 The Client Assignment Workflow (core business process)')
bullets([
    "For each unit in a project, staff select the client who owns it, label the "
    "unit (e.g. 'Unit A1'), set the contract amount, set a status (active/"
    "inactive), set a progress percentage, and set start/handover dates.",
    "This creates an 'assignment' — the link between a client and a project unit "
    "— which becomes the basis for everything the client sees in their portal "
    "(documents, media, payments, updates for that unit).",
])

h2('4.7 Payments & Financial Tracking')
bullets([
    "A list of all recorded payments, showing the client, project, amount, "
    "status (pending/paid/overdue), payment method, reference, due date, and paid "
    "date.",
    "Staff record new payments against a specific client assignment; the system "
    "automatically updates that assignment's 'amount paid' and 'outstanding "
    "balance' figures.",
    "Staff can also directly adjust an assignment's contract amount, amount paid, "
    "outstanding balance, and next payment due date/amount if corrections are "
    "needed — all changes are logged.",
])

h2('4.8 Document Management')
bullets([
    "A central library of every document uploaded for any project or client, "
    "searchable and filterable.",
    "Staff upload documents (PDFs, Word, Excel, etc.), attach them to a project "
    "and/or a specific client assignment, give them a title, category (e.g. "
    "Contract, Specification), version, and notes.",
    "Each document can be marked visible or hidden from the client, and archived "
    "if no longer current.",
    "Documents can be edited (replace the file, change details) or deleted.",
])

h2('4.9 Media Management')
bullets([
    "A media library of all photos and videos across all projects, filterable by "
    "type, project, and gallery section.",
    "Staff can upload images or videos (up to 200MB) or link to externally-hosted "
    "media, attach a thumbnail, give it a title and notes, mark it as featured, "
    "set its display order, choose which gallery section it belongs to, and "
    "control whether clients can see it.",
    "Media can be edited or deleted at any time.",
])

h2('4.10 Project Updates')
bullets([
    "Staff publish progress updates for any project: a title, description, the "
    "date, and a progress percentage (0–100%).",
    "Each update can be made visible or hidden from clients. Visible updates "
    "appear instantly in the relevant clients' portals.",
])

h2('4.11 Content Management System (CMS)')
doc.add_paragraph(
    "The CMS lets staff change almost everything visitors see on the public "
    "website — text, images, and settings — without needing a developer. Changes "
    "appear on the live site immediately. Sections include:"
)
bullets([
    "Homepage — every section (hero banner, featured projects, services overview, "
    "'why choose us' points, process steps, testimonials, blog preview, portal "
    "call-to-action, final call-to-action, and company stats).",
    "Pages — the text/layout content for Services, About, Portfolio, Blog, FAQ, "
    "Careers, Contact, Partnerships, and Testimonials pages.",
    "About Page — hero text, team members (add/edit/remove, with photo and bio), "
    "and the company history timeline (add/edit/remove entries).",
    "Services — create and edit each service offering, including its benefits, "
    "process steps, and which projects it's linked to.",
    "Blog — write, edit, and publish blog articles, mark a post as featured, set "
    "categories.",
    "Testimonials — add/edit client reviews, ratings, and which ones are "
    "featured.",
    "FAQ — add/edit questions and answers, organised by category.",
    "Careers — create/edit/close job postings (applications come in through the "
    "public Careers page).",
    "Portfolio — edit the text and metrics shown on the Portfolio page.",
    "Image/video uploads (up to 200MB) are supported throughout the CMS.",
])

h2('4.12 Forms & Enquiries Inbox')
bullets([
    "A single place to view everything submitted through the public website's "
    "forms: contact messages, project enquiries, career applications, partnership "
    "enquiries, support tickets, and newsletter signups.",
])

h2('4.13 Notifications')
bullets([
    "A live alert centre summarising anything needing staff attention: pending "
    "consultations, pending KYC reviews, overdue payments, open support tickets, "
    "and new contact submissions — the same alerts shown on the dashboard, in one "
    "place.",
])

h2('4.14 Settings')
bullets([
    "Company-wide details used throughout the site: company name, tagline, phone "
    "numbers, email addresses, business hours, and address.",
    "Text-size scaling controls for the public website, the client portal, and the "
    "admin area, so Rockmould can make text larger or smaller across the whole "
    "platform if needed.",
])

h2('4.15 Staff Roles & Permissions Management')
bullets([
    "A list of all staff/admin accounts (separate from clients), filterable by "
    "role.",
    "Admins can change a staff member's role or activate/deactivate their "
    "account.",
    "A permissions matrix shows, at a glance, which roles can access which parts "
    "of the admin area.",
])

h2('4.16 Activity Log')
bullets([
    "A complete, time-stamped audit trail of significant actions across the "
    "system — who did what and when (e.g. a project was created, a KYC request "
    "was approved, a payment was recorded, a role was changed). This provides "
    "accountability and helps with troubleshooting.",
])


# ============================================================
# SECTION 5: KEY WORKFLOWS
# ============================================================
h1('5. How It All Fits Together — Key Workflows')

h2('5.1 New Client Onboarding')
bullets([
    "A prospective client submits a verification (KYC) form on the portal login "
    "page, including which project they're interested in.",
    "The request appears in the Admin Dashboard's KYC review queue.",
    "Staff review the details and either approve or reject the request.",
    "On approval, the system automatically creates the client's portal account.",
    "The client can now log in using the email + one-time-code process.",
    "Staff then assign the client to a project/unit (Section 4.6.1), after which "
    "the client sees that project, its documents, media, updates, and payment "
    "balance in their portal.",
])

h2('5.2 Project Sale & Financial Tracking')
bullets([
    "Staff create the project in the Admin Dashboard with all its details and "
    "media.",
    "Staff assign a unit within that project to a client, recording the agreed "
    "contract amount.",
    "As the client pays (in instalments or in full), staff record each payment "
    "against that assignment.",
    "The system automatically keeps the 'amount paid' and 'outstanding balance' "
    "up to date — both staff and the client can see this in real time.",
    "Staff post progress updates as construction/delivery proceeds; the client "
    "sees these in their portal.",
])

h2('5.3 Consultation Booking')
bullets([
    "A visitor picks an available date/time on the public Consultation page and "
    "submits their details.",
    "A booking is created with a unique reference; confirmation emails go to both "
    "the visitor and Rockmould.",
    "Staff review and manage the booking from the Admin Dashboard — assigning a "
    "team member, adding a meeting link, and updating its status — with the "
    "client notified by email of any changes.",
])

h2('5.4 Website Content Updates')
bullets([
    "Staff log in to the Admin Dashboard and open the relevant CMS section "
    "(homepage, a specific page, blog, services, etc.).",
    "Staff edit text, replace images/videos, or add new entries (e.g. a new blog "
    "post, team member, or FAQ).",
    "Changes are saved and appear on the public website immediately — no "
    "developer involvement or website redeployment is required.",
])


# ============================================================
# SUMMARY
# ============================================================
h1('Summary')

doc.add_paragraph(
    "The Rockmould platform is, in effect, three connected systems sharing one "
    "set of data: a marketing/public website that attracts and informs visitors, a "
    "Client Portal that gives investors a transparent, real-time view of their "
    "projects and finances, and an Admin Dashboard that gives Rockmould staff full "
    "control over projects, clients, payments, content, and day-to-day operations "
    "— all protected by role-based access and a complete activity log for "
    "accountability."
)

doc.save(r"C:\xampp\htdocs\rockmould-legacy-main\reports\Rockmould_System_Overview.docx")
print("done")
